"""Point-in-time check of the Huazheng panel against Huazheng's own documents (contract 27.1).

    python -I -X utf8 vintage_sources.py

Inputs (local, not committed): the two PDFs fetched from chindices.com on 2026-10-28 and kept in
data/raw/huazheng_official_20261028/, the netdisk quarterly panel and the older netdisk copy. The tracking report is
a scanned brochure without a text layer; its 2018-10-31 grade distribution and the table of the 15 largest
downgrades were transcribed from the page images (page 2 and 3) and are typed in below. Writes vintage_sources.json
(LF): source URLs, read date, file hashes, the transcribed figures and the panel's figures for the same date and firms.
"""

from __future__ import annotations

import collections
import datetime as dt
import hashlib
import json
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OFFICIAL = ROOT / "data/raw/huazheng_official_20261028"
PANEL = ROOT / "data/raw/huazheng_panel_netdisk_20261026/huazheng_quarterly_2009_2025.xlsx"
OLD = ROOT / "data/raw/huazheng_panel_netdisk_20261013/huazheng_2009_2022.xlsx"
GRADES = ["C", "CC", "CCC", "B", "BB", "BBB", "A", "AA", "AAA"]

SOURCES = {
    "faq": {"url": "https://www.chindices.com/esg-faq.html", "read": "2026-10-28",
            "quotes": ["华证股票ESG评级数据按季度更新，评级日期分别为每年0131、0430、0731和1031",
                       "更新后的评级数据一般会在评级日期的20个交易日内发布",
                       "华证A股ESG评级数据最早从2009年开始，其中2009-2018年为1.0体系评级结果，2019年至今为2.0体系评级结果"]},
    "tracking_report": {"url": "https://www.chindices.com/files/上海华证指数ESG跟踪分析服务.pdf", "read": "2026-10-28",
                        "file": "上海华证指数ESG跟踪分析服务.pdf",
                        "note": "brochure '华证指数 ESG 跟踪'; page 1: '评价回溯至 2009 年 1 月'; pages 2-3 describe 2018-10-31"},
    "methodology_v2_1": {"url": "https://www.chindices.com/files/华证ESG评级方法论V2.1.pdf", "read": "2026-10-28",
                         "file": "华证ESG评级方法论V2.1.pdf",
                         "versions": {"2022.03": "V1.0 first public methodology", "2022.05": "V1.1",
                                      "2022.11": "V2.0 full update of the framework",
                                      "2024.05": "V2.1 industry adjustment; weights aligned with the exchanges' guidelines"}},
}
# Tracking report, page 2: grade distribution of 2018-10-31 (count of listed firms).
REPORT_2018_10 = {"AAA": 107, "AA": 542, "A": 654, "BBB": 1604, "BB": 452, "B": 148, "CCC": 40, "CC": 8, "C": 1}
# Page 3, left table: the 15 largest downgrades of the period (rating 2018-10-31, rating of the previous period).
REPORT_DOWNGRADES = {"300152": ("CC", "BBB"), "603505": ("CCC", "A"), "600397": ("CCC", "BBB"), "002291": ("CCC", "BBB"),
                     "002317": ("CCC", "BBB"), "000971": ("CCC", "BBB"), "300392": ("CCC", "BBB"), "300087": ("B", "BBB"),
                     "000513": ("B", "BBB"), "002001": ("BBB", "AA"), "002484": ("B", "BBB"), "300664": ("B", "BBB"),
                     "600200": ("BBB", "AA"), "002340": ("BBB", "AA"), "002029": ("B", "A")}
# Page 3, chart: score of 上峰水泥 (000672) on 2016-07-31.
REPORT_SCORE_000672_2016_07 = 89.85


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def mean_notch(counts: dict) -> float:
    n = sum(counts.values())
    return sum((GRADES.index(g) + 1) * c for g, c in counts.items()) / n


def main() -> None:
    import openpyxl

    out: dict = {"sources": SOURCES}
    for s in SOURCES.values():
        if "file" in s:
            p = OFFICIAL / s["file"]
            s["sha256"], s["bytes"] = sha(p), p.stat().st_size
    rows = list(openpyxl.load_workbook(PANEL, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    d10, d07, d1607 = dt.datetime(2018, 10, 31), dt.datetime(2018, 7, 31), dt.datetime(2016, 7, 31)
    panel = {(f"{r[0]:06d}", r[1]): (r[3], r[4]) for r in rows[1:]}
    dist = collections.Counter(g for (c, d), (g, s) in panel.items() if d == d10)
    old_rows = list(openpyxl.load_workbook(OLD, read_only=True, data_only=True)["季度数据"].iter_rows(values_only=True))
    j = [str(h)[:10] for h in old_rows[0]].index("2018-12-31")            # the old copy labels 2018-10-31 as 2018-12-31
    old_dist = collections.Counter(r[j] for r in old_rows[1:] if r[j] in GRADES)
    firms = {}
    for code, (now, before) in REPORT_DOWNGRADES.items():
        a, b = panel.get((code, d10)), panel.get((code, d07))
        firms[code] = {"report": [now, before], "report_change": GRADES.index(now) - GRADES.index(before),
                       "panel": [a[0], b[0]] if a and b else None,
                       "panel_change": GRADES.index(a[0]) - GRADES.index(b[0]) if a and b else None}
    out["2018-10-31"] = {
        "report": {"firms": sum(REPORT_2018_10.values()), "counts": REPORT_2018_10, "mean_notch": mean_notch(REPORT_2018_10)},
        "panel": {"firms": sum(dist.values()), "counts": {g: dist.get(g, 0) for g in GRADES[::-1]}, "mean_notch": mean_notch(dist)},
        "old_copy": {"firms": sum(old_dist.values()), "counts": {g: old_dist.get(g, 0) for g in GRADES[::-1]},
                     "mean_notch": mean_notch(old_dist)},
        "report_downgrades_in_panel": firms,
        "score_000672_2016_07_31": {"report": REPORT_SCORE_000672_2016_07, "panel": panel[("000672", d1607)][1]},
    }
    text = json.dumps(out, ensure_ascii=False, indent=2)
    (HERE / "vintage_sources.json").write_bytes((text + "\n").encode("utf-8"))
    print(text)


if __name__ == "__main__":
    main()
